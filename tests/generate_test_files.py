import os
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import docx

def create_test_files():
    docs_dir = Path("data/docs")
    docs_dir.mkdir(parents=True, exist_ok=True)

    # 1. Plain Text file
    txt_path = docs_dir / "cloud_security_guidelines.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("""
Cloud Security and Data Governance Guidelines 2026

1. Authentication and MFA
All corporate systems require Multi-Factor Authentication (MFA) via hardware security keys or authenticator applications. SMS-based 2FA is deprecated.

2. API Keys and Secrets Management
Hardcoding API tokens in source code is strictly prohibited. All secrets must be stored in HashiCorp Vault or AWS Secrets Manager. Secrets found in Git commits will result in immediate token revocation and an automated incident alert.

3. Database Encryption
All customer relational databases must have AES-256 encryption enabled at rest. Backups are replicated across two geographic cloud zones with 30-day immutability retention.
""".strip())
    print(f"Created: {txt_path}")

    # 2. Microsoft Word Document
    docx_path = docs_dir / "employee_wellness_policy.docx"
    doc = docx.Document()
    doc.add_heading("Corporate Wellness and Mental Health Policy 2026", level=1)
    doc.add_paragraph("Enterprise Global Technologies provides comprehensive wellness stipends to encourage work-life balance.")
    doc.add_heading("1. Fitness and Gym Reimbursement", level=2)
    doc.add_paragraph("Employees can claim up to ₹18,000 annually for gym memberships, yoga studios, or home fitness equipment.")
    doc.add_heading("2. Mental Health Counseling", level=2)
    doc.add_paragraph("Every employee is entitled to 12 free confidential therapy sessions per year through our telehealth partner Lyra Health.")
    doc.save(docx_path)
    print(f"Created: {docx_path}")

    # 3. Image File with text to test OCR
    img_path = docs_dir / "office_catering_menu.png"
    img = Image.new("RGB", (700, 300), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    
    draw.text((40, 30), "CORPORATE CATERING AND SNACK POLICY", fill=(20, 20, 20))
    draw.text((40, 70), "Complimentary cafeteria breakfast is served daily from 8:30 AM to 10:30 AM.", fill=(40, 40, 40))
    draw.text((40, 110), "Friday Team Lunches are subsidized 100% by the department food budget.", fill=(40, 40, 40))
    draw.text((40, 150), "Special dietary meal requests must be submitted 24 hours in advance.", fill=(40, 40, 40))
    draw.text((40, 190), "Office Coffee Bar: Unlimited espresso, cold brew, and matcha available 24/7.", fill=(40, 40, 40))
    
    img.save(img_path)
    print(f"Created image for OCR: {img_path}")

if __name__ == "__main__":
    create_test_files()
